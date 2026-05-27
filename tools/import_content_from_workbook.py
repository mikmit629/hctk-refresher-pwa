#!/usr/bin/env python3
"""Import HCTK activity content into the refresher PWA."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
APP_DIR = SCRIPT_DIR.parent
CODING_DIR = APP_DIR.parent
DEFAULT_WORKBOOK = CODING_DIR / "hctk_refresher_activities_26_draft.xlsx"
DEFAULT_INTERVENTION_APP = APP_DIR / "app.js"
DEFAULT_CONTROL_APP = APP_DIR / "control" / "app.js"
ACTIVITY_COUNT = 26
SHEET_NAME = "Refresher Activities"
CHOICE_KEYS = ["choice_a", "choice_b", "choice_c", "choice_d"]
TARGETS = {
    "intervention": {
        "path": DEFAULT_INTERVENTION_APP,
        "id_prefix": "content",
    },
    "control": {
        "path": DEFAULT_CONTROL_APP,
        "id_prefix": "control-content",
    },
}
REQUIRED_HEADERS = [
    "activity_id / calendar sequence",
    "format / card type",
    "title / calendar summary",
    "detail / calendar description",
    "media / card asset",
    "choice_a",
    "choice_b",
    "choice_c",
    "choice_d",
    "correct_choice",
    "rationale / card feedback",
    "citation_numbers",
    "citation_references",
]


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Transfer HCTK workbook activity rows into the PWA CONTENT_BANK array."
    )
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--target", choices=sorted(TARGETS), default="intervention")
    parser.add_argument("--intervention-app", type=Path, default=DEFAULT_INTERVENTION_APP)
    parser.add_argument("--control-app", type=Path, default=DEFAULT_CONTROL_APP)
    parser.add_argument("--app", type=Path, help="Override the target app.js path.")
    parser.add_argument("--dry-run", action="store_true", help="Validate and report without writing app files.")
    args = parser.parse_args()

    workbook_path = args.workbook.expanduser().resolve()
    if not workbook_path.exists():
        raise SystemExit(f"Workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, data_only=False)
    target = TARGETS[args.target]
    rows = read_refresher_sheet(workbook, target["id_prefix"])
    version = content_version(rows, args.target)

    path = target_app_path(args).expanduser().resolve()
    if not path.exists():
        raise SystemExit(f"App file not found: {path}")
    source = path.read_text(encoding="utf-8")
    next_text = replace_app_content(source, rows, version)
    changed = next_text != source
    if args.dry_run:
        print(f"Would update {args.target}: {path} ({len(rows)} activities, version {version}, changed={changed})")
    elif changed:
        path.write_text(next_text, encoding="utf-8")
        print(f"Updated {args.target}: {path} ({len(rows)} activities, version {version})")
    else:
        print(f"No content changes for {args.target}: {path} ({len(rows)} activities, version {version})")


def target_app_path(args: argparse.Namespace) -> Path:
    if args.app:
        return args.app
    if args.target == "control":
        return args.control_app
    return args.intervention_app


def read_refresher_sheet(workbook: Any, id_prefix: str) -> list[dict[str, Any]]:
    if SHEET_NAME not in workbook.sheetnames:
        raise SystemExit(f"Required sheet missing: {SHEET_NAME}")

    sheet = workbook[SHEET_NAME]
    headers = {
        normalize_header(cell.value): index + 1
        for index, cell in enumerate(sheet[1])
        if normalize_header(cell.value)
    }
    missing = [name for name in REQUIRED_HEADERS if name not in headers]
    if missing:
        raise SystemExit(f"{SHEET_NAME} missing required columns: {', '.join(missing)}")

    rows: list[dict[str, Any]] = []
    for row_number in range(2, sheet.max_row + 1):
        activity_id = cell_text(sheet, headers, row_number, "activity_id / calendar sequence")
        if not activity_id and row_is_empty(sheet, row_number, len(headers)):
            continue
        sequence = parse_activity_sequence(activity_id, row_number)
        content_format = cell_text(sheet, headers, row_number, "format / card type")
        validate_format(content_format, row_number)

        choices = [
            cell_text(sheet, headers, row_number, name) for name in CHOICE_KEYS
        ]
        correct_choice = cell_text(sheet, headers, row_number, "correct_choice")
        rationale = cell_text(sheet, headers, row_number, "rationale / card feedback")
        correct_index = validate_card_fields(content_format, choices, correct_choice, rationale, row_number)
        if content_format == "Question":
            rationale = clean_rationale(rationale)
        media_path = cell_text(sheet, headers, row_number, "media / card asset")
        media_type = infer_media_type(media_path) if media_path else ""
        rows.append({
            "id": f"{id_prefix}-{sequence:03d}",
            "activityId": activity_id,
            "sequence": sequence,
            "format": content_format,
            "title": cell_text(sheet, headers, row_number, "title / calendar summary"),
            "prompt": cell_text(sheet, headers, row_number, "detail / calendar description"),
            "mediaType": media_type,
            "mediaPath": media_path,
            "mediaAlt": "",
            "image": media_path if media_type and media_type != "link" else "",
            "choices": [choice for choice in choices if choice],
            "correctIndex": correct_index,
            "rationale": rationale,
            "citationNumbers": cell_text(sheet, headers, row_number, "citation_numbers"),
            "citationReferences": cell_text(sheet, headers, row_number, "citation_references"),
        })

    rows.sort(key=lambda item: item["sequence"])
    if len(rows) != ACTIVITY_COUNT:
        raise SystemExit(f"{SHEET_NAME} must contain exactly {ACTIVITY_COUNT} activities; found {len(rows)}")
    if len({row["id"] for row in rows}) != ACTIVITY_COUNT:
        raise SystemExit("Generated content ids must be unique")
    if [row["sequence"] for row in rows] != list(range(1, ACTIVITY_COUNT + 1)):
        raise SystemExit(f"{SHEET_NAME} must contain activity ids 1/26 through 26/26")
    return rows


def normalize_header(value: Any) -> str:
    return str(value or "").strip().lower()


def cell_text(sheet: Any, headers: dict[str, int], row_number: int, name: str) -> str:
    column = headers.get(name)
    if not column:
        return ""
    value = sheet.cell(row=row_number, column=column).value
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def row_is_empty(sheet: Any, row_number: int, max_column: int) -> bool:
    for column in range(1, max_column + 1):
        value = sheet.cell(row=row_number, column=column).value
        if value is not None and str(value).strip():
            return False
    return True


def parse_activity_sequence(value: str, row_number: int) -> int:
    match = re.fullmatch(r"(\d+)\s*/\s*26", value)
    if not match:
        raise SystemExit(f"{SHEET_NAME} row {row_number}: activity_id / calendar sequence must use N/26")
    sequence = int(match.group(1))
    if sequence < 1 or sequence > ACTIVITY_COUNT:
        raise SystemExit(f"{SHEET_NAME} row {row_number}: activity sequence is outside 1/26 through 26/26")
    return sequence


def validate_format(value: str, row_number: int) -> None:
    if value not in {"Fact", "Question"}:
        raise SystemExit(f"{SHEET_NAME} row {row_number}: format / card type must be Fact or Question")


def validate_card_fields(
    content_format: str,
    choices: list[str],
    correct_choice: str,
    rationale: str,
    row_number: int,
) -> int:
    if content_format == "Fact":
        if any(choices) or correct_choice or rationale:
            raise SystemExit(f"{SHEET_NAME} row {row_number}: Fact rows must not include choices, correct choice, or rationale")
        return 0

    if not all(choices):
        raise SystemExit(f"{SHEET_NAME} row {row_number}: Question rows must include choices A through D")
    if correct_choice not in {"a", "b", "c", "d"}:
        raise SystemExit(f"{SHEET_NAME} row {row_number}: correct_choice must be lowercase a, b, c, or d")
    if not rationale:
        raise SystemExit(f"{SHEET_NAME} row {row_number}: Question rows must include rationale / card feedback")
    return {"a": 0, "b": 1, "c": 2, "d": 3}[correct_choice]


def clean_rationale(value: str) -> str:
    lines = [line.strip() for line in value.splitlines() if line.strip()]
    has_choice_list = sum(1 for line in lines if rationale_choice_line(line)) >= 2
    cleaned: list[str] = []

    for line in lines:
        if has_choice_list and re.match(r"^Correct answer:\s*[A-D]\b", line, flags=re.IGNORECASE):
            continue
        match = rationale_choice_line(line)
        if match:
            label = match.group(1).upper()
            detail = re.sub(r"^Correct\.\s*", "", match.group(2), flags=re.IGNORECASE)
            cleaned.append(f"{label}: {detail}")
            continue
        cleaned.append(re.sub(r"^(Correct answer:\s*[A-D]\.\s*)Correct\.\s*", r"\1", line, flags=re.IGNORECASE))

    return "\n".join(cleaned)


def rationale_choice_line(value: str) -> re.Match[str] | None:
    return re.match(r"^([A-D])[:.)]?\s+(.+)$", value, flags=re.IGNORECASE)


def infer_media_type(path: str) -> str:
    text = str(path or "").strip().lower()
    if text.startswith(("http://", "https://")):
        return "link"
    if re.search(r"\.(mp4|webm|mov|m4v)(\?|#|$)", text):
        return "video"
    if re.search(r"\.gif(\?|#|$)", text):
        return "gif"
    return "image"


def content_version(rows: list[dict[str, Any]], target: str) -> str:
    payload = json.dumps(
        {target: rows},
        ensure_ascii=True,
        sort_keys=True,
        separators=(",", ":"),
    )
    return f"spreadsheet-{hashlib.sha256(payload.encode()).hexdigest()[:12]}"


def replace_app_content(source: str, rows: list[dict[str, Any]], version: str) -> str:
    content_bank = format_content_bank(rows)
    next_source, version_count = re.subn(
        r"window\.HCTK_CONTENT_BANK_VERSION\s*=\s*[^;]*;",
        f"window.HCTK_CONTENT_BANK_VERSION = '{version}';",
        source,
        count=1,
    )
    if version_count != 1:
        raise SystemExit("Could not find HCTK_CONTENT_BANK_VERSION assignment to replace")
    pattern = re.compile(r"window\.HCTK_CONTENT_BANK = \[[\s\S]*?\n\];")
    next_source, count = pattern.subn(lambda _match: content_bank, next_source, count=1)
    if count != 1:
        raise SystemExit("Could not find HCTK_CONTENT_BANK block to replace")
    return next_source


def format_content_bank(rows: list[dict[str, Any]]) -> str:
    parts = ["window.HCTK_CONTENT_BANK = ["]
    for index, item in enumerate(rows):
        parts.extend([
            "  {",
            f"    id: {js_string(item['id'])},",
            f"    activityId: {js_string(item['activityId'])},",
            f"    sequence: {int(item['sequence'])},",
            f"    format: {js_string(item['format'])},",
            f"    title: {js_string(item['title'])},",
            f"    prompt: {js_string(item['prompt'])},",
            f"    mediaType: {js_string(item['mediaType'])},",
            f"    mediaPath: {js_string(item['mediaPath'])},",
            f"    mediaAlt: {js_string(item['mediaAlt'])},",
            f"    image: {js_string(item['image'])},",
            "    choices: [",
        ])
        for choice_index, choice in enumerate(item["choices"]):
            comma = "," if choice_index < len(item["choices"]) - 1 else ""
            parts.append(f"      {js_string(choice)}{comma}")
        parts.extend([
            "    ],",
            f"    correctIndex: {int(item['correctIndex'])},",
            f"    rationale: {js_string(item['rationale'])},",
            f"    citationNumbers: {js_string(item['citationNumbers'])},",
            f"    citationReferences: {js_string(item['citationReferences'])}",
            "  }" + ("," if index < len(rows) - 1 else ""),
        ])
    parts.append("];")
    return "\n".join(parts)


def js_string(value: Any) -> str:
    text = str(value or "")
    escaped = (
        text
        .replace("\\", "\\\\")
        .replace("\n", "\\n")
        .replace("'", "\\'")
    )
    return f"'{escaped}'"


if __name__ == "__main__":
    main()
